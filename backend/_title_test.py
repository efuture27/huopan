# -*- coding: utf-8 -*-
"""离线验证：导出 Excel 标题格式 + 导出文件可被导入逻辑重新解析"""
import sys, io
sys.path.insert(0, r"E:\WorkBuddyData\WorkBuddy\2026-08-23-11-00-35\ai-huopan-demo\backend")

from openpyxl import load_workbook
from app.services.excel import generate_plan_excel, parse_product_excel

lines = []
def log(s):
    lines.append(str(s))

# 模拟生成货盘（含 1px PNG data URL 图片）
png_b64 = ("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
items = [
    {"sku": "A001", "name": "示例针织开衫", "merchant": "某某服饰有限公司",
     "category": "上衣", "cost_price": 80, "sale_price": 199,
     "image_url": "data:image/png;base64," + png_b64},
    {"sku": "A002", "name": "示例半身裙", "merchant": "某某服饰有限公司",
     "category": "裙子", "cost_price": 60, "sale_price": 149, "image_url": ""},
]
ai = {"descriptions": {"A001": "测试描述"}, "selling_points": {"A001": "测试卖点"}}
buf = generate_plan_excel(items, ai, "测试货盘")
data = buf.getvalue()
log("export ok: %d bytes" % len(data))

# 验证结构
wb = load_workbook(io.BytesIO(data))
ws = wb.active
log("title A1: %r" % ws["A1"].value)
log("merged: %s" % [str(r) for r in ws.merged_cells.ranges])
log("header row2: %s" % [c.value for c in ws[2]])
log("row3 first 3: %s" % [c.value for c in ws[3]][:3])
log("freeze: %s" % ws.freeze_panes)
log("images: %d" % len(ws._images))
log("header font bold=%s color=%s fill=%s" % (
    ws["A2"].font.bold, ws["A2"].font.color.rgb if ws["A2"].font.color else None,
    ws["A2"].fill.fgColor.rgb))
log("title font size=%s color=%s fill=%s" % (
    ws["A1"].font.sz, ws["A1"].font.color.rgb if ws["A1"].font.color else None,
    ws["A1"].fill.fgColor.rgb))

# 闭环：导出文件重新导入
prods = parse_product_excel(data, filename="测试货盘.xlsx")
log("reimport: %d products" % len(prods))
if prods:
    log("first: sku=%s name=%s merchant=%s cost=%s" % (
        prods[0]["sku"], prods[0]["name"], prods[0]["merchant"], prods[0]["cost_price"]))
    log("img recovered: %s" % bool(prods[0]["image_url"]))

with open(r"E:\WorkBuddyData\WorkBuddy\2026-08-23-11-00-35\ai-huopan-demo\backend\_title_test_out.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("DONE")
