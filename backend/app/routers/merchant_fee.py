import io
from fastapi import APIRouter, Depends, UploadFile, File, HTTPException
from sqlalchemy.orm import Session
from ..database import get_db
from .. import models, schemas, security
from openpyxl import load_workbook

router = APIRouter(prefix="/api/merchant-fee", tags=["merchant-fee"])

# 导入表表头别名：第一列商家名，第二列运费说明
_MERCHANT_ALIAS = ("商家", "商家名称", "供应商", "供货商", "厂家", "merchant", "supplier")
_FEE_ALIAS = ("运费说明", "偏远运费", "偏远地区运费", "偏远地区加收运费", "偏远费", "运费", "规则", "fee_note", "note", "备注")


def _parse_fee_excel(data: bytes):
    """解析两列导入表：第一列商家名、第二列运费说明。返回 [(merchant, fee_note), ...]"""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    # 定位表头行：需「同一行」同时命中商家类与运费类表头，
    # 避免标题行（如「商家运费导入模板」）被误判为表头。
    header_idx = None
    m_col = f_col = None
    cand = None  # 只有商家列、尚无运费列的候选行（兜底）
    for i, row in enumerate(rows[:8]):
        row_m = row_f = None
        for j, cell in enumerate(row):
            s = str(cell or "").strip()
            if not s:
                continue
            if row_m is None and any(a in s for a in _MERCHANT_ALIAS):
                row_m = j
            elif row_f is None and any(a in s for a in _FEE_ALIAS):
                row_f = j
        if row_m is not None and row_f is not None:
            m_col, f_col, header_idx = row_m, row_f, i
            break
        if row_m is not None and cand is None:
            cand = (row_m, i)
    if header_idx is None:
        if cand is None:
            raise ValueError("未找到「商家」列表头，请使用模板格式（第一列商家名称，第二列运费说明）")
        m_col, header_idx = cand
    if f_col is None:
        f_col = m_col + 1  # 约定第二列

    out = []
    for row in rows[header_idx + 1:]:
        if m_col >= len(row):
            continue
        merchant = str(row[m_col] or "").strip()
        if not merchant or merchant.startswith("（") and merchant.endswith("）") and len(merchant) <= 4:
            continue  # 跳过空行与占位说明
        fee = ""
        if f_col < len(row) and row[f_col] is not None:
            fee = str(row[f_col]).strip()
        out.append((merchant, fee))
    return out


@router.get("/list")
def list_fees(
    db: Session = Depends(get_db),
    user: str = Depends(security.get_current_user),
):
    """全部商家运费配置。"""
    rows = db.query(models.MerchantFee).order_by(models.MerchantFee.merchant).all()
    return {
        "rows": [{"merchant": r.merchant, "fee_note": r.fee_note or ""} for r in rows],
    }


def _upsert(db: Session, merchant: str, fee_note: str):
    merchant = (merchant or "").strip()
    if not merchant:
        raise HTTPException(400, "商家名不能为空")
    rec = db.query(models.MerchantFee).filter(models.MerchantFee.merchant == merchant).first()
    if rec:
        rec.fee_note = fee_note or ""
    else:
        db.add(models.MerchantFee(merchant=merchant, fee_note=fee_note or ""))
    db.commit()


@router.post("/save")
def save_fee(
    req: schemas.MerchantFeeItem,
    db: Session = Depends(get_db),
    user: str = Depends(security.get_current_user),
):
    """保存单个商家运费说明（存在则更新）。"""
    _upsert(db, req.merchant, req.fee_note)
    return {"ok": True, "merchant": req.merchant}


@router.post("/batch")
def batch_save(
    req: schemas.MerchantFeeBatch,
    db: Session = Depends(get_db),
    user: str = Depends(security.get_current_user),
):
    """批量保存商家运费说明。"""
    saved = 0
    for it in req.items:
        m = (it.merchant or "").strip()
        if not m:
            continue
        _upsert(db, m, it.fee_note)
        saved += 1
    return {"ok": True, "saved": saved}


@router.post("/upload")
def upload_fees(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: str = Depends(security.get_current_user),
):
    """上传两列导入表（商家名称 / 运费说明）批量写入，已有商家更新。"""
    data = file.file.read()
    try:
        pairs = _parse_fee_excel(data)
    except ValueError as e:
        raise HTTPException(400, str(e))
    saved = 0
    for merchant, fee in pairs:
        _upsert(db, merchant, fee)
        saved += 1
    return {"ok": True, "saved": saved, "merchants": [m for m, _ in pairs]}


@router.delete("/{merchant}")
def delete_fee(
    merchant: str,
    db: Session = Depends(get_db),
    user: str = Depends(security.get_current_user),
):
    """删除某商家的运费配置（商品库商家不受影响）。"""
    rec = db.query(models.MerchantFee).filter(models.MerchantFee.merchant == merchant).first()
    if rec:
        db.delete(rec)
        db.commit()
        return {"ok": True, "deleted": merchant}
    return {"ok": True, "deleted": None}
