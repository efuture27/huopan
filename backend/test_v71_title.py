# -*- coding: utf-8 -*-
"""v7.1 冒烟：主标题写 Excel 首行（公域紫字/私域黑字、表头下移、freeze A3）+ 无标题回归 + 重新下载。"""
import io
import json
import urllib.parse
import urllib.request
from openpyxl import load_workbook

BASE = "http://127.0.0.1:8000"
PASSED, FAILED = [], []


def check(name, cond, detail=""):
    (PASSED if cond else FAILED).append(f"{name} {detail}")
    print(("  PASS  " if cond else "  FAIL  ") + name + ("" if cond else "  " + str(detail)))


def http(method, path, token=None, data=None, form=False):
    req = urllib.request.Request(BASE + path, method=method)
    if token:
        req.add_header("Authorization", "Bearer " + token)
    if data is not None:
        if form:
            req.data = urllib.parse.urlencode(data).encode()
            req.add_header("Content-Type", "application/x-www-form-urlencoded")
        else:
            req.data = json.dumps(data).encode()
            req.add_header("Content-Type", "application/json")
    with urllib.request.urlopen(req, timeout=60) as r:
        body = r.read()
        ct = r.headers.get("Content-Type", "")
        return r.status, (json.loads(body) if "json" in ct else body)


import urllib.parse  # noqa: E402

ITEMS = [
    {"sku": "T001", "name": "冒烟商品A", "cost_price": 50, "shipping_fee": 5, "sale_price": 129,
     "spec": "均码", "product_type": "袋泡茶", "origin": "云南昆明", "excluded_regions": "新疆西藏"},
    {"sku": "T002", "name": "冒烟商品B", "cost_price": 30, "shipping_fee": 3, "sale_price": 89},
]

# ---------- 登录 ----------
st, login = http("POST", "/api/auth/login", data={"username": "admin", "password": "admin123"}, form=True)
TOKEN = login["access_token"]
check("登录", st == 200 and TOKEN)

# ---------- 公域带主标题 ----------
st, x1 = http("POST", "/api/plan/generate", TOKEN, {
    "name": "v71冒烟-公域", "title": "9月公域达人货盘", "plan_type": "public",
    "commission_cut": 0, "cost_mode": "daifa", "tax_rate": 0, "items": ITEMS})
wb = load_workbook(io.BytesIO(x1))
ws = wb.active
check("公域 A1 主标题", ws["A1"].value == "9月公域达人货盘", repr(ws["A1"].value))
check("公域 A1:H1 合并", "A1:H1" in [str(m) for m in ws.merged_cells.ranges], [str(m) for m in ws.merged_cells.ranges])
check("公域 A2=序号", ws["A2"].value == "序号", repr(ws["A2"].value))
check("公域 表头紫底", (ws["B2"].fill.fgColor.rgb or "")[-6:] == "8B5CF6", ws["B2"].fill.fgColor.rgb)
check("公域 freeze A3", ws.freeze_panes == "A3", ws.freeze_panes)
check("公域 B3 商品名", ws["B3"].value == "冒烟商品A", repr(ws["B3"].value))
check("公域 G3 销售价", float(ws["G3"].value) == 129.0, ws["G3"].value)
check("公域 H3 佣金57.4%", ws["H3"].value == "57.4%", repr(ws["H3"].value))
st, plans = http("GET", "/api/plan/list", TOKEN)
pid_pub = plans[0]["id"]

# ---------- 私域带主标题 ----------
st, x2 = http("POST", "/api/plan/generate", TOKEN, {
    "name": "v71冒烟-私域", "title": "鑫多多 9 月私域报价", "plan_type": "private",
    "tax_rate": 0.09, "items": ITEMS})
ws = load_workbook(io.BytesIO(x2)).active
check("私域 A1 主标题", ws["A1"].value == "鑫多多 9 月私域报价", repr(ws["A1"].value))
check("私域 A1:V1 合并", "A1:V1" in [str(m) for m in ws.merged_cells.ranges], [str(m) for m in ws.merged_cells.ranges])
check("私域 A2=序号", ws["A2"].value == "序号", repr(ws["A2"].value))
check("私域 B2 黄底", (ws["B2"].fill.fgColor.rgb or "")[-6:] == "FFFF00", ws["B2"].fill.fgColor.rgb)
check("私域 freeze A3", ws.freeze_panes == "A3", ws.freeze_panes)
check("私域 H3 直播价", float(ws["H3"].value) == 129.0, ws["H3"].value)
check("私域 I3 一件代发55", float(ws["I3"].value) == 55.0, ws["I3"].value)
check("私域 J3 含税59.95", abs(float(ws["J3"].value) - 59.95) < 1e-6, ws["J3"].value)
check("私域 K3 毛利率57.4%", ws["K3"].value == "57.4%", repr(ws["K3"].value))
check("私域 B3 产品类型", ws["B3"].value == "袋泡茶", repr(ws["B3"].value))

# ---------- 无标题回归（v7.0 行为不变） ----------
st, x3 = http("POST", "/api/plan/generate", TOKEN, {
    "name": "v71冒烟-无标题", "plan_type": "public", "commission_cut": 0,
    "cost_mode": "daifa", "tax_rate": 0, "items": ITEMS})
ws = load_workbook(io.BytesIO(x3)).active
check("无标题 A1=序号", ws["A1"].value == "序号", repr(ws["A1"].value))
check("无标题 freeze A2", ws.freeze_panes == "A2", ws.freeze_panes)
check("无标题 无合并行", all("A1:" not in str(m) for m in ws.merged_cells.ranges))

# ---------- 重新下载带主标题 ----------
st, plans = http("GET", "/api/plan/list", TOKEN)
pid_nofile = [p for p in plans if p["name"] == "v71冒烟-私域"][0]["id"]
st, x4 = http("GET", f"/api/plan/{pid_pub}/download", TOKEN)
ws = load_workbook(io.BytesIO(x4)).active
check("重新下载 A1 主标题", ws["A1"].value == "9月公域达人货盘", repr(ws["A1"].value))
st, x5 = http("GET", f"/api/plan/{pid_nofile}/download", TOKEN)
ws = load_workbook(io.BytesIO(x5)).active
check("重新下载私域 A1 主标题", ws["A1"].value == "鑫多多 9 月私域报价", repr(ws["A1"].value))

# ---------- 清理测试货盘 ----------
names = {"v71冒烟-公域", "v71冒烟-私域", "v71冒烟-无标题"}
st, plans = http("GET", "/api/plan/list", TOKEN)
for p in [p for p in plans if p["name"] in names]:
    http("DELETE", f"/api/plan/{p['id']}", TOKEN)
st, plans = http("GET", "/api/plan/list", TOKEN)
check("清理完成", not [p for p in plans if p["name"] in names], [p["name"] for p in plans])

print(f"\n==== v7.1 冒烟结果: {len(PASSED)} PASS / {len(FAILED)} FAIL ====")
if FAILED:
    for f in FAILED:
        print("  FAILED:", f)
    raise SystemExit(1)
